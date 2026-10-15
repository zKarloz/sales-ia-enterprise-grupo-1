"""Generación de archivos Excel XLSX para reportes F12."""

import csv

from datetime import datetime
from decimal import Decimal
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.schemas.report import (
    SalesReportResponse,
    InventoryReportResponse,
    CustomerReportResponse,
    AnalyticsReportResponse,
)
from app.services.report_export_csv import (
    SALES_COLUMNS,
    build_sales_csv,
    INVENTORY_PRODUCT_COLUMNS,
    INVENTORY_MOVEMENT_COLUMNS,
    build_inventory_products_csv,
    build_inventory_movements_csv,
    CUSTOMER_COLUMNS,
    build_customers_csv,
)
from app.services.report_export_analytics_csv import (
    ANALYTICS_COLUMNS,
    build_analytics_csv_datasets,
)


HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="17365D",
)

HEADER_FONT = Font(
    bold=True,
    color="FFFFFF",
)

MONEY_FORMAT = '#,##0.00'
DATE_FORMAT = "yyyy-mm-dd hh:mm:ss"


def _style_sheet(ws):
    """Aplica formato básico empresarial a la hoja."""

    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    ws.row_dimensions[1].height = 28

    for column_cells in ws.columns:
        letter = get_column_letter(
            column_cells[0].column
        )

        max_length = max(
            len(str(cell.value or ""))
            for cell in column_cells
        )

        ws.column_dimensions[letter].width = min(
            max(max_length + 3, 14),
            38,
        )

    ws.auto_filter.ref = ws.dimensions


def build_sales_xlsx(
    report: SalesReportResponse,
) -> bytes:
    """Genera un Excel del reporte de ventas."""

    csv_bytes = build_sales_csv(report)

    rows = csv.reader(
        StringIO(
            csv_bytes.decode("utf-8-sig"),
            newline="",
        )
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas"

    for row_number, row in enumerate(rows, start=1):

        if row_number == 1:
            ws.append(row)
            continue

        converted = []

        for index, value in enumerate(row):

            # Identificadores numéricos.
            if index in (0, 2, 4, 6):
                converted.append(
                    int(value) if value else None
                )

            # Importes monetarios.
            elif index in (8, 9, 10, 11):
                converted.append(
                    Decimal(value) if value else None
                )

            # Fecha de venta.
            elif index == 1 and value:
                parsed = datetime.fromisoformat(value)

                # Excel almacena fechas sin zona horaria.
                converted.append(
                    parsed.replace(tzinfo=None)
                )

            else:
                # Los textos proceden del generador CSV,
                # que neutraliza fórmulas no confiables.
                converted.append(value)

        ws.append(converted)

    _style_sheet(ws)

    for row in ws.iter_rows(min_row=2):
        row[1].number_format = DATE_FORMAT

        for column_index in (8, 9, 10, 11):
            row[column_index].number_format = MONEY_FORMAT

    output = BytesIO()
    wb.save(output)

    return output.getvalue()

def build_inventory_xlsx(
    report: InventoryReportResponse,
) -> bytes:
    """Genera Excel con productos y movimientos."""

    wb = Workbook()

    products_sheet = wb.active
    products_sheet.title = "Productos"

    movements_sheet = wb.create_sheet("Movimientos")

    datasets = [
        (
            products_sheet,
            build_inventory_products_csv(report),
            INVENTORY_PRODUCT_COLUMNS,
        ),
        (
            movements_sheet,
            build_inventory_movements_csv(report),
            INVENTORY_MOVEMENT_COLUMNS,
        ),
    ]

    for ws, content, columns in datasets:
        rows = csv.reader(
            StringIO(
                content.decode("utf-8-sig"),
                newline="",
            )
        )

        ws.append(columns)

        for row_number, row in enumerate(rows):
            if row_number == 0:
                continue

            converted = []

            for index, value in enumerate(row):
                if ws.title == "Productos":
                    if index in (0, 3, 5):
                        value = int(value) if value else None
                    elif index in (6, 7):
                        value = Decimal(value) if value else None
                    elif index in (8, 9):
                        value = value == "True"
                else:
                    if index in (0, 2, 5, 6, 7, 8):
                        value = int(value) if value else None
                    elif index == 1 and value:
                        value = datetime.fromisoformat(
                            value
                        ).replace(tzinfo=None)

                converted.append(value)

            ws.append(converted)

        _style_sheet(ws)

    for row in products_sheet.iter_rows(min_row=2):
        row[6].number_format = MONEY_FORMAT
        row[7].number_format = MONEY_FORMAT

    for row in movements_sheet.iter_rows(min_row=2):
        row[1].number_format = DATE_FORMAT

    output = BytesIO()
    wb.save(output)
    wb.close()

    return output.getvalue()

def _append_csv_to_sheet(
    ws,
    content: bytes,
    *,
    integer_columns=(),
    decimal_columns=(),
    date_columns=(),
    boolean_columns=(),
):
    """Convierte un CSV validado a una hoja Excel tipada.

    Los índices de columna comienzan en cero.
    """

    rows = csv.reader(
        StringIO(
            content.decode("utf-8-sig"),
            newline="",
        )
    )

    for row_number, row in enumerate(rows):
        if row_number == 0:
            ws.append(row)
            continue

        converted = []

        for index, value in enumerate(row):
            if not value:
                converted.append(None)
                continue

            if index in integer_columns:
                converted.append(int(value))

            elif index in decimal_columns:
                # openpyxl almacena números con la
                # precisión numérica habitual de Excel.
                converted.append(float(Decimal(value)))

            elif index in date_columns:
                parsed = datetime.fromisoformat(value)
                converted.append(
                    parsed.replace(tzinfo=None)
                )

            elif index in boolean_columns:
                converted.append(value == "True")

            else:
                # Conserva como texto los identificadores
                # documentales y las cadenas protegidas.
                converted.append(value)

        ws.append(converted)

    _style_sheet(ws)

    for row in ws.iter_rows(min_row=2):
        for index in decimal_columns:
            row[index].number_format = MONEY_FORMAT

        for index in date_columns:
            row[index].number_format = DATE_FORMAT


def _save_workbook(wb: Workbook) -> bytes:
    """Serializa y cierra un libro Excel."""

    output = BytesIO()

    try:
        wb.save(output)
        return output.getvalue()
    finally:
        wb.close()

def build_customers_xlsx(
    report: CustomerReportResponse,
) -> bytes:
    """Genera el reporte de clientes en Excel."""

    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"

    _append_csv_to_sheet(
        ws,
        build_customers_csv(report),
        integer_columns=(0, 7),
        decimal_columns=(8, 9),
        date_columns=(10,),
        boolean_columns=(6,),
    )

    return _save_workbook(wb)

ANALYTICS_SHEET_CONFIG = {
    "kpis": {
        "title": "KPIs",
        "integers": (1, 2, 3),
        "decimals": (0, 4, 5, 6),
        "dates": (),
    },
    "periods": {
        "title": "Periodos",
        "integers": (3,),
        "decimals": (2,),
        "dates": (0,),
    },
    "products": {
        "title": "Productos",
        "integers": (0, 2),
        "decimals": (3, 4),
        "dates": (),
    },
    "categories": {
        "title": "Categorias",
        "integers": (0, 2),
        "decimals": (3, 4),
        "dates": (),
    },
    "sellers": {
        "title": "Vendedores",
        "integers": (0, 2),
        "decimals": (3, 4),
        "dates": (),
    },
    "distributions": {
        "title": "Distribuciones",
        "integers": (3,),
        "decimals": (1, 2, 4),
        "dates": (),
    },
    "comparisons": {
        "title": "Comparaciones",
        "integers": (),
        "decimals": (0, 1, 2, 3),
        "dates": (),
    },
}


def build_analytics_xlsx(
    report: AnalyticsReportResponse,
) -> bytes:
    """Genera las siete hojas del reporte Analytics."""

    datasets = build_analytics_csv_datasets(report)

    wb = Workbook()
    wb.remove(wb.active)

    for name, config in ANALYTICS_SHEET_CONFIG.items():
        ws = wb.create_sheet(config["title"])

        _append_csv_to_sheet(
            ws,
            datasets[name],
            integer_columns=config["integers"],
            decimal_columns=config["decimals"],
            date_columns=config["dates"],
        )

    return _save_workbook(wb)
