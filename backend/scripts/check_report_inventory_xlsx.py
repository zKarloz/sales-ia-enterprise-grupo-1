"""Pruebas unitarias del XLSX de inventario F12.

No consulta ni modifica Supabase.
"""

import sys
import unittest

from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from zipfile import is_zipfile

from openpyxl import load_workbook

sys.path.insert(
    0,
    str(Path(__file__).resolve().parents[1]),
)

from app.schemas.report import (
    InventoryReportFilters,
    InventoryReportResponse,
    InventoryReportSummary,
    InventoryProductItem,
    InventoryMovementItem,
)

from app.services.report_export_csv import (
    INVENTORY_PRODUCT_COLUMNS,
    INVENTORY_MOVEMENT_COLUMNS,
)

from app.services.report_export_xlsx import (
    build_inventory_xlsx,
)


def make_product(**changes):
    data = {
        "product_id": 2,
        "sku": "PROD-002",
        "product_name": "Mouse Inalámbrico",
        "category_id": 2,
        "category_name": "Periféricos",
        "stock": 78,
        "unit_price": Decimal("25.00"),
        "reference_value": Decimal("1950.00"),
        "is_active": True,
        "low_stock": False,
    }

    data.update(changes)

    return InventoryProductItem.model_construct(**data)


def make_movement(**changes):
    data = {
        "movement_id": 9,
        "created_at": datetime(
            2026, 10, 9, 16, 39,
            tzinfo=timezone.utc,
        ),
        "product_id": 2,
        "product_name": "Mouse Inalámbrico",
        "movement_type": "IN",
        "quantity": 2,
        "stock_before": 78,
        "stock_after": 80,
        "supplier_id": None,
        "reason": "Reposición de prueba",
    }

    data.update(changes)

    return InventoryMovementItem.model_construct(**data)


def make_report(products=None, movements=None):
    return InventoryReportResponse(
        filters_applied=InventoryReportFilters(
            date_from="2026-10-01",
            date_to="2026-10-31",
        ),
        summary=InventoryReportSummary(),
        products=products if products is not None else [],
        movements=movements if movements is not None else [],
    )


def open_excel(report):
    content = build_inventory_xlsx(report)

    workbook = load_workbook(
        BytesIO(content),
        data_only=False,
    )

    return content, workbook


class InventoryXlsxTests(unittest.TestCase):

    def test_valid_excel(self):
        content, workbook = open_excel(make_report())

        try:
            self.assertIsInstance(content, bytes)
            self.assertTrue(is_zipfile(BytesIO(content)))
        finally:
            workbook.close()

    def test_two_sheets(self):
        _, workbook = open_excel(make_report())

        try:
            self.assertEqual(
                workbook.sheetnames,
                ["Productos", "Movimientos"],
            )
        finally:
            workbook.close()

    def test_headers(self):
        _, workbook = open_excel(make_report())

        try:
            products = workbook["Productos"]
            movements = workbook["Movimientos"]

            self.assertEqual(
                [cell.value for cell in products[1]],
                INVENTORY_PRODUCT_COLUMNS,
            )

            self.assertEqual(
                [cell.value for cell in movements[1]],
                INVENTORY_MOVEMENT_COLUMNS,
            )
        finally:
            workbook.close()

    def test_empty_report(self):
        _, workbook = open_excel(make_report())

        try:
            self.assertEqual(
                workbook["Productos"].max_row,
                1,
            )

            self.assertEqual(
                workbook["Movimientos"].max_row,
                1,
            )
        finally:
            workbook.close()

    def test_product_values(self):
        _, workbook = open_excel(
            make_report(products=[make_product()])
        )

        try:
            sheet = workbook["Productos"]

            self.assertEqual(sheet["A2"].value, 2)
            self.assertEqual(sheet["C2"].value, "Mouse Inalámbrico")
            self.assertEqual(sheet["F2"].value, 78)
            self.assertEqual(sheet["G2"].value, 25)
            self.assertEqual(sheet["H2"].value, 1950)
        finally:
            workbook.close()

    def test_product_numeric_types(self):
        _, workbook = open_excel(
            make_report(products=[make_product()])
        )

        try:
            sheet = workbook["Productos"]

            for column in ("F", "G", "H"):
                with self.subTest(column=column):
                    self.assertEqual(
                        sheet[f"{column}2"].data_type,
                        "n",
                    )
        finally:
            workbook.close()

    def test_boolean_values(self):
        _, workbook = open_excel(
            make_report(
                products=[
                    make_product(
                        is_active=True,
                        low_stock=False,
                    )
                ]
            )
        )

        try:
            sheet = workbook["Productos"]

            self.assertIs(sheet["I2"].value, True)
            self.assertIs(sheet["J2"].value, False)
        finally:
            workbook.close()

    def test_movement_values(self):
        _, workbook = open_excel(
            make_report(movements=[make_movement()])
        )

        try:
            sheet = workbook["Movimientos"]

            self.assertEqual(sheet["A2"].value, 9)
            self.assertEqual(sheet["E2"].value, "IN")
            self.assertEqual(sheet["F2"].value, 2)
            self.assertEqual(sheet["G2"].value, 78)
            self.assertEqual(sheet["H2"].value, 80)
        finally:
            workbook.close()

    def test_null_values(self):
        _, workbook = open_excel(
            make_report(
                movements=[
                    make_movement(
                        stock_before=None,
                        stock_after=None,
                        supplier_id=None,
                    )
                ]
            )
        )

        try:
            sheet = workbook["Movimientos"]

            self.assertIsNone(sheet["G2"].value)
            self.assertIsNone(sheet["H2"].value)
            self.assertIsNone(sheet["I2"].value)
        finally:
            workbook.close()

    def test_movement_date(self):
        _, workbook = open_excel(
            make_report(movements=[make_movement()])
        )

        try:
            sheet = workbook["Movimientos"]

            self.assertEqual(
                sheet["B2"].value,
                datetime(2026, 10, 9, 16, 39),
            )

            self.assertEqual(
                sheet["B2"].number_format,
                "yyyy-mm-dd hh:mm:ss",
            )
        finally:
            workbook.close()

    def test_formatting(self):
        _, workbook = open_excel(
            make_report(
                products=[make_product()],
                movements=[make_movement()],
            )
        )

        try:
            for sheet in workbook.worksheets:
                with self.subTest(sheet=sheet.title):
                    self.assertEqual(
                        sheet.freeze_panes,
                        "A2",
                    )

                    self.assertIsNotNone(
                        sheet.auto_filter.ref
                    )

                    self.assertTrue(
                        sheet["A1"].font.bold
                    )
        finally:
            workbook.close()

    def test_formula_protection(self):
        _, workbook = open_excel(
            make_report(
                products=[
                    make_product(
                        product_name="=SUM(A1:A2)",
                    )
                ]
            )
        )

        try:
            cell = workbook["Productos"]["C2"]

            self.assertNotEqual(cell.data_type, "f")
            self.assertTrue(cell.value.startswith("'"))
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
