"""Pruebas HTTP del XLSX de inventario F12."""

import sys
import unittest

from datetime import date
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from main import app
from app.api.dependencies.auth import get_current_user
from app.core.database import get_db
from app.core.roles import (
    ROLE_ADMIN,
    ROLE_MANAGER,
    ROLE_ANALYST,
    ROLE_SELLER,
    ROLE_WAREHOUSE,
)


URL = (
    "/api/reports/inventory/export/xlsx"
    "?date_from=2026-10-01&date_to=2026-10-31"
)

XLSX_MIME = (
    "application/vnd.openxmlformats-officedocument."
    "spreadsheetml.sheet"
)


def sample_xlsx():
    workbook = Workbook()
    products = workbook.active
    products.title = "Productos"
    products.append(["ID Producto", "Stock"])
    products.append([1, 10])

    movements = workbook.create_sheet("Movimientos")
    movements.append(["ID Movimiento", "Cantidad"])
    movements.append([1, 5])

    output = BytesIO()
    workbook.save(output)
    workbook.close()

    return output.getvalue()


class InventoryXlsxApiTests(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app)

    @classmethod
    def tearDownClass(cls):
        cls.client.close()

    def setUp(self):
        app.dependency_overrides.clear()
        app.dependency_overrides[get_db] = lambda: object()

    def tearDown(self):
        app.dependency_overrides.clear()

    def authenticate(self, role):
        app.dependency_overrides[get_current_user] = (
            lambda: SimpleNamespace(id=1, role=role)
        )

    def request_export(self, url=URL):
        with (
            patch(
                "app.api.routes.reports.build_inventory_report",
                return_value=object(),
            ) as report_service,
            patch(
                "app.api.routes.reports.build_inventory_xlsx",
                return_value=sample_xlsx(),
            ) as xlsx_service,
        ):
            response = self.client.get(url)

        return response, report_service, xlsx_service

    def test_allowed_roles(self):
        for role in (ROLE_ADMIN, ROLE_MANAGER):
            with self.subTest(role=role):
                self.authenticate(role)
                response, _, _ = self.request_export()
                self.assertEqual(response.status_code, 200)

    def test_unauthenticated(self):
        response = self.client.get(URL)
        self.assertEqual(response.status_code, 401)

    def test_forbidden_roles(self):
        for role in (
            ROLE_ANALYST,
            ROLE_SELLER,
            ROLE_WAREHOUSE,
        ):
            with self.subTest(role=role):
                self.authenticate(role)
                response = self.client.get(URL)
                self.assertEqual(response.status_code, 403)

    def test_headers(self):
        self.authenticate(ROLE_MANAGER)
        response, _, _ = self.request_export()

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.headers["content-type"],
            XLSX_MIME,
        )
        self.assertEqual(
            response.headers["cache-control"],
            "no-store",
        )
        self.assertIn(
            "inventory_report_2026-10-01_2026-10-31.xlsx",
            response.headers["content-disposition"],
        )

    def test_filters_forwarded(self):
        self.authenticate(ROLE_MANAGER)

        response, service, _ = self.request_export(
            URL
            + "&category_id=2"
            + "&stock_threshold=8"
            + "&include_inactive=true"
        )

        self.assertEqual(response.status_code, 200)

        filters = service.call_args.kwargs["filters"]

        self.assertEqual(filters.category_id, 2)
        self.assertEqual(filters.stock_threshold, 8)
        self.assertTrue(filters.include_inactive)

    def test_invalid_dates(self):
        self.authenticate(ROLE_MANAGER)

        response = self.client.get(
            "/api/reports/inventory/export/xlsx"
            "?date_from=2026-10-31"
            "&date_to=2026-10-01"
        )

        self.assertEqual(response.status_code, 422)

    def test_invalid_filters(self):
        self.authenticate(ROLE_MANAGER)

        for parameter in ("category_id=0", "stock_threshold=-1"):
            with self.subTest(parameter=parameter):
                response = self.client.get(
                    URL + "&" + parameter
                )
                self.assertEqual(response.status_code, 422)

    def test_excel_content(self):
        self.authenticate(ROLE_MANAGER)

        response, _, _ = self.request_export()

        workbook = load_workbook(
            BytesIO(response.content),
            data_only=True,
        )

        try:
            self.assertEqual(
                workbook.sheetnames,
                ["Productos", "Movimientos"],
            )
            self.assertEqual(
                workbook["Productos"]["B2"].value,
                10,
            )
            self.assertEqual(
                workbook["Movimientos"]["B2"].value,
                5,
            )
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
