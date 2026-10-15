"""Pruebas XLSX de clientes y Analytics F12."""

import sys
import unittest

from datetime import datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(
    0,
    str(Path(__file__).resolve().parents[1]),
)

from check_report_customers_csv import (
    make_customer,
    make_report as make_customers_report,
)
from check_report_analytics_csv import (
    make_report as make_analytics_report,
)

from app.services.report_export_xlsx import (
    build_customers_xlsx,
    build_analytics_xlsx,
)


def open_workbook(content):
    return load_workbook(
        BytesIO(content),
        data_only=False,
    )


class NewXlsxTests(unittest.TestCase):

    def test_customers_sheet(self):
        report = make_customers_report([
            make_customer()
        ])

        wb = open_workbook(
            build_customers_xlsx(report)
        )

        try:
            self.assertEqual(
                wb.sheetnames,
                ["Clientes"],
            )
            self.assertEqual(
                wb["Clientes"].max_column,
                11,
            )
        finally:
            wb.close()

    def test_customer_numeric_values(self):
        report = make_customers_report([
            make_customer()
        ])

        wb = open_workbook(
            build_customers_xlsx(report)
        )

        try:
            ws = wb["Clientes"]

            self.assertEqual(ws["A2"].value, 6)
            self.assertEqual(ws["H2"].value, 3)
            self.assertAlmostEqual(
                ws["I2"].value,
                1250.50,
            )
            self.assertEqual(
                ws["I2"].data_type,
                "n",
            )
            self.assertIs(
                ws["G2"].value,
                True,
            )
        finally:
            wb.close()

    def test_customer_document_text(self):
        report = make_customers_report([
            make_customer(
                document_number="00123456789",
                phone="001234567",
            )
        ])

        wb = open_workbook(
            build_customers_xlsx(report)
        )

        try:
            ws = wb["Clientes"]

            self.assertEqual(
                ws["D2"].value,
                "00123456789",
            )
            self.assertEqual(
                ws["F2"].value,
                "001234567",
            )
        finally:
            wb.close()

    def test_customer_date_and_empty(self):
        report = make_customers_report([
            make_customer(),
            make_customer(
                customer_id=7,
                last_purchase_at=None,
            ),
        ])

        wb = open_workbook(
            build_customers_xlsx(report)
        )

        try:
            ws = wb["Clientes"]

            self.assertIsInstance(
                ws["K2"].value,
                datetime,
            )
            self.assertIsNone(
                ws["K3"].value,
            )
        finally:
            wb.close()

    def test_customers_empty_report(self):
        wb = open_workbook(
            build_customers_xlsx(
                make_customers_report()
            )
        )

        try:
            self.assertEqual(
                wb["Clientes"].max_row,
                1,
            )
        finally:
            wb.close()

    def test_customer_formula_protection(self):
        report = make_customers_report([
            make_customer(
                customer_name="=SUM(A1:A2)",
            )
        ])

        wb = open_workbook(
            build_customers_xlsx(report)
        )

        try:
            cell = wb["Clientes"]["B2"]

            self.assertNotEqual(
                cell.data_type,
                "f",
            )
            self.assertTrue(
                cell.value.startswith("'"),
            )
        finally:
            wb.close()

    def test_analytics_seven_sheets(self):
        wb = open_workbook(
            build_analytics_xlsx(
                make_analytics_report()
            )
        )

        try:
            self.assertEqual(
                wb.sheetnames,
                [
                    "KPIs",
                    "Periodos",
                    "Productos",
                    "Categorias",
                    "Vendedores",
                    "Distribuciones",
                    "Comparaciones",
                ],
            )
        finally:
            wb.close()

    def test_analytics_kpis(self):
        wb = open_workbook(
            build_analytics_xlsx(
                make_analytics_report()
            )
        )

        try:
            ws = wb["KPIs"]

            self.assertAlmostEqual(
                ws["A2"].value,
                12537.45,
            )
            self.assertEqual(
                ws["B2"].value,
                11,
            )
            self.assertEqual(
                ws["D2"].value,
                50,
            )
        finally:
            wb.close()

    def test_analytics_numeric_and_date(self):
        wb = open_workbook(
            build_analytics_xlsx(
                make_analytics_report()
            )
        )

        try:
            products = wb["Productos"]
            periods = wb["Periodos"]

            self.assertEqual(
                products["D2"].data_type,
                "n",
            )
            self.assertAlmostEqual(
                products["D2"].value,
                10000.00,
            )
            self.assertIsInstance(
                periods["A2"].value,
                datetime,
            )
        finally:
            wb.close()

    def test_negative_comparison(self):
        wb = open_workbook(
            build_analytics_xlsx(
                make_analytics_report(
                    negative_change=True
                )
            )
        )

        try:
            ws = wb["Comparaciones"]

            self.assertEqual(
                ws["C2"].value,
                -100,
            )
            self.assertEqual(
                ws["D2"].value,
                -10,
            )
            self.assertEqual(
                ws["C2"].data_type,
                "n",
            )
        finally:
            wb.close()

    def test_empty_analytics_sections(self):
        wb = open_workbook(
            build_analytics_xlsx(
                make_analytics_report(
                    comparisons=False,
                    empty_lists=True,
                )
            )
        )

        try:
            self.assertEqual(
                wb["Comparaciones"].max_row,
                1,
            )
            self.assertEqual(
                wb["Productos"].max_row,
                1,
            )
            self.assertEqual(
                len(wb.sheetnames),
                7,
            )
        finally:
            wb.close()

    def test_sheet_formatting(self):
        wb = open_workbook(
            build_analytics_xlsx(
                make_analytics_report()
            )
        )

        try:
            for ws in wb.worksheets:
                with self.subTest(sheet=ws.title):
                    self.assertEqual(
                        ws.freeze_panes,
                        "A2",
                    )
                    self.assertTrue(
                        ws["A1"].font.bold,
                    )
                    self.assertIsNotNone(
                        ws.auto_filter.ref,
                    )
        finally:
            wb.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
