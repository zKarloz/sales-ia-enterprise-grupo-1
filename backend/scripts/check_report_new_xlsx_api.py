"""Pruebas HTTP y RBAC de XLSX: clientes y Analytics F12.

No consulta ni modifica Supabase.
"""

import sys
import unittest

from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from main import app
from app.api.dependencies.auth import get_current_user
from app.core.database import get_db
from app.core.roles import (
    ROLE_ADMIN,
    ROLE_MANAGER,
    ROLE_ANALYST,
    ROLE_SELLER,
    ROLE_WAREHOUSE,
)


BASE_URL = "/api/reports"
DATES = "?date_from=2026-10-01&date_to=2026-10-31"

XLSX_MIME = (
    "application/vnd.openxmlformats-officedocument."
    "spreadsheetml.sheet"
)

CASES = {
    "customers": {
        "url": BASE_URL + "/customers/export/xlsx" + DATES,
        "report_service": "build_customers_report",
        "xlsx_service": "build_customers_xlsx",
        "sheet": "Clientes",
    },
    "analytics": {
        "url": BASE_URL + "/analytics/export/xlsx" + DATES,
        "report_service": "build_analytics_report",
        "xlsx_service": "build_analytics_xlsx",
        "sheet": "KPIs",
    },
}


def sample_xlsx(sheet_name):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(["ID", "Valor"])
    sheet.append([1, 12537.45])

    output = BytesIO()
    workbook.save(output)
    workbook.close()

    return output.getvalue()


class NewXlsxApiTests(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app)

    @classmethod
    def tearDownClass(cls):
        cls.client.close()

    def setUp(self):
        app.dependency_overrides.clear()
        app.dependency_overrides[get_db] = lambda: object()

    def tearDown(self):
        app.dependency_overrides.clear()

    def authenticate(self, role):
        app.dependency_overrides[get_current_user] = (
            lambda: SimpleNamespace(id=1, role=role)
        )

    def request_export(self, case, url=None):
        with (
            patch(
                "app.api.routes.reports."
                + case["report_service"],
                return_value=object(),
            ) as report_mock,
            patch(
                "app.api.routes.reports."
                + case["xlsx_service"],
                return_value=sample_xlsx(case["sheet"]),
            ) as xlsx_mock,
        ):
            response = self.client.get(
                url or case["url"]
            )

        return response, report_mock, xlsx_mock

    def test_allowed_roles(self):
        for name, case in CASES.items():
            for role in (ROLE_ADMIN, ROLE_MANAGER):
                with self.subTest(name=name, role=role):
                    self.authenticate(role)
                    response, _, _ = self.request_export(case)
                    self.assertEqual(response.status_code, 200)

    def test_unauthenticated(self):
        for name, case in CASES.items():
            with self.subTest(name=name):
                response = self.client.get(case["url"])
                self.assertEqual(response.status_code, 401)

    def test_forbidden_roles(self):
        for name, case in CASES.items():
            for role in (
                ROLE_ANALYST,
                ROLE_SELLER,
                ROLE_WAREHOUSE,
            ):
                with self.subTest(name=name, role=role):
                    self.authenticate(role)
                    response = self.client.get(case["url"])
                    self.assertEqual(response.status_code, 403)

    def test_headers(self):
        for name, case in CASES.items():
            with self.subTest(name=name):
                self.authenticate(ROLE_MANAGER)
                response, _, _ = self.request_export(case)

                self.assertEqual(response.status_code, 200)
                self.assertEqual(
                    response.headers["content-type"],
                    XLSX_MIME,
                )
                self.assertEqual(
                    response.headers["cache-control"],
                    "no-store",
                )
                self.assertIn(
                    f"{name}_report_2026-10-01_2026-10-31.xlsx",
                    response.headers["content-disposition"],
                )
                self.assertIn(
                    "attachment",
                    response.headers["content-disposition"],
                )

    def test_invalid_dates(self):
        for name, case in CASES.items():
            with self.subTest(name=name):
                self.authenticate(ROLE_MANAGER)

                url = (
                    BASE_URL
                    + f"/{name}/export/xlsx"
                    + "?date_from=2026-10-31"
                    + "&date_to=2026-10-01"
                )

                with patch(
                    "app.api.routes.reports."
                    + case["report_service"]
                ) as service:
                    response = self.client.get(url)

                self.assertEqual(response.status_code, 422)
                service.assert_not_called()

    def test_invalid_filters(self):
        for name, case in CASES.items():
            params = [
                "branch_id=0",
                "seller_id=-1",
            ]
            if name == "analytics":
                params.append("category_id=0")

            for parameter in params:
                with self.subTest(
                    name=name,
                    parameter=parameter,
                ):
                    self.authenticate(ROLE_MANAGER)
                    response = self.client.get(
                        case["url"] + "&" + parameter
                    )
                    self.assertEqual(response.status_code, 422)

    def test_filters_forwarded(self):
        for name, case in CASES.items():
            with self.subTest(name=name):
                self.authenticate(ROLE_MANAGER)

                extra = (
                    "&include_inactive=true"
                    if name == "customers"
                    else "&category_id=3"
                )

                response, service, xlsx_service = (
                    self.request_export(
                        case,
                        case["url"]
                        + "&branch_id=1&seller_id=2"
                        + extra,
                    )
                )

                self.assertEqual(response.status_code, 200)

                filters = service.call_args.kwargs["filters"]
                self.assertEqual(filters.branch_id, 1)
                self.assertEqual(filters.seller_id, 2)

                if name == "customers":
                    self.assertTrue(filters.include_inactive)
                else:
                    self.assertEqual(filters.category_id, 3)

                xlsx_service.assert_called_once()

    def test_excel_readable(self):
        for name, case in CASES.items():
            with self.subTest(name=name):
                self.authenticate(ROLE_MANAGER)
                response, _, _ = self.request_export(case)

                workbook = load_workbook(
                    BytesIO(response.content),
                    data_only=True,
                )

                try:
                    self.assertIn(
                        case["sheet"],
                        workbook.sheetnames,
                    )
                    sheet = workbook[case["sheet"]]
                    self.assertEqual(sheet["A2"].value, 1)
                    self.assertEqual(
                        sheet["B2"].value,
                        12537.45,
                    )
                finally:
                    workbook.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
