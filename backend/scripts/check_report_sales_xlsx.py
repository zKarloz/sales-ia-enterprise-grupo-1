"""Pruebas unitarias del exportador XLSX de ventas F12.

No requiere conexión a Supabase.
"""

import sys
import unittest

from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from zipfile import is_zipfile

from openpyxl import load_workbook

sys.path.insert(
    0,
    str(Path(__file__).resolve().parents[1]),
)

from app.schemas.report import SalesReportResponse
from app.services.report_export_csv import SALES_COLUMNS
from app.services.report_export_xlsx import build_sales_xlsx


def make_sale(**changes):
    data = {
        "sale_id": 1,
        "created_at": datetime(
            2026, 10, 9, 15, 30,
            tzinfo=timezone.utc,
        ),
        "customer_id": 6,
        "customer_name": "Empresa Norte S.A.C.",
        "seller_id": 2,
        "seller_name": "Vendedor Prueba",
        "branch_id": 1,
        "branch_name": "Principal",
        "subtotal_amount": Decimal("1000.00"),
        "discount_amount": Decimal("50.00"),
        "tax_amount": Decimal("180.00"),
        "total_amount": Decimal("1130.00"),
        "payment_method": "TARJETA",
    }

    data.update(changes)
    return type("SaleStub", (), data)()


def make_report(items=None):
    return SalesReportResponse.model_construct(
        items=items if items is not None else []
    )


def open_excel(report):
    content = build_sales_xlsx(report)
    workbook = load_workbook(
        BytesIO(content),
        data_only=False,
    )
    return content, workbook


class SalesXlsxTests(unittest.TestCase):

    def test_valid_excel_bytes(self):
        content, workbook = open_excel(make_report())

        self.assertIsInstance(content, bytes)
        self.assertTrue(is_zipfile(BytesIO(content)))
        self.assertEqual(workbook.sheetnames, ["Ventas"])

        workbook.close()

    def test_headers(self):
        _, workbook = open_excel(make_report())
        sheet = workbook["Ventas"]

        headers = [
            cell.value
            for cell in sheet[1]
        ]

        self.assertEqual(headers, SALES_COLUMNS)
        self.assertEqual(sheet.max_column, 13)

        workbook.close()

    def test_empty_report(self):
        _, workbook = open_excel(make_report())
        sheet = workbook["Ventas"]

        self.assertEqual(sheet.max_row, 1)

        workbook.close()

    def test_sale_values(self):
        _, workbook = open_excel(
            make_report([make_sale()])
        )

        sheet = workbook["Ventas"]

        self.assertEqual(sheet["A2"].value, 1)
        self.assertEqual(sheet["D2"].value, "Empresa Norte S.A.C.")
        self.assertEqual(sheet["L2"].value, 1130)
        self.assertEqual(sheet["M2"].value, "TARJETA")

        workbook.close()

    def test_numeric_amounts(self):
        _, workbook = open_excel(
            make_report([
                make_sale(
                    subtotal_amount=Decimal("100.50"),
                    discount_amount=Decimal("10.25"),
                    tax_amount=Decimal("2.50"),
                    total_amount=Decimal("92.75"),
                )
            ])
        )

        sheet = workbook["Ventas"]

        for column in ("I", "J", "K", "L"):
            with self.subTest(column=column):
                cell = sheet[f"{column}2"]

                self.assertEqual(cell.data_type, "n")
                self.assertIsInstance(
                    cell.value,
                    (int, float),
                )

        self.assertAlmostEqual(
            sheet["L2"].value,
            92.75,
        )

        workbook.close()

    def test_date_value(self):
        _, workbook = open_excel(
            make_report([make_sale()])
        )

        sheet = workbook["Ventas"]

        self.assertEqual(
            sheet["B2"].value,
            datetime(2026, 10, 9, 15, 30),
        )

        self.assertEqual(
            sheet["B2"].number_format,
            "yyyy-mm-dd hh:mm:ss",
        )

        workbook.close()

    def test_excel_formatting(self):
        _, workbook = open_excel(
            make_report([make_sale()])
        )

        sheet = workbook["Ventas"]

        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertIsNotNone(sheet.auto_filter.ref)
        self.assertFalse(
            sheet.sheet_view.showGridLines
        )
        self.assertTrue(sheet["A1"].font.bold)

        workbook.close()

    def test_special_characters(self):
        name = 'Compañía "Norte", S.A.C.'

        _, workbook = open_excel(
            make_report([
                make_sale(customer_name=name)
            ])
        )

        self.assertEqual(
            workbook["Ventas"]["D2"].value,
            name,
        )

        workbook.close()

    def test_formula_protection(self):
        _, workbook = open_excel(
            make_report([
                make_sale(
                    customer_name="=SUM(A1:A2)",
                )
            ])
        )

        cell = workbook["Ventas"]["D2"]

        self.assertNotEqual(cell.data_type, "f")
        self.assertTrue(cell.value.startswith("'"))

        workbook.close()

    def test_multiple_sales(self):
        _, workbook = open_excel(
            make_report([
                make_sale(sale_id=1),
                make_sale(sale_id=2),
                make_sale(sale_id=3),
            ])
        )

        sheet = workbook["Ventas"]

        self.assertEqual(sheet.max_row, 4)
        self.assertEqual(sheet["A4"].value, 3)

        workbook.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
