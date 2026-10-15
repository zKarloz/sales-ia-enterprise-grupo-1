"""Pruebas HTTP y RBAC del XLSX de ventas F12.

No consulta ni modifica Supabase.
"""

import sys
import unittest

from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

sys.path.insert(
    0,
    str(Path(__file__).resolve().parents[1]),
)

from fastapi.testclient import TestClient

from main import app
from app.api.dependencies.auth import get_current_user
from app.core.database import get_db
from app.core.roles import (
    ROLE_ADMIN,
    ROLE_ANALYST,
    ROLE_MANAGER,
    ROLE_SELLER,
    ROLE_WAREHOUSE,
)


URL = (
    "/api/reports/sales/export/xlsx"
    "?date_from=2026-10-01"
    "&date_to=2026-10-31"
)

XLSX_MIME = (
    "application/vnd.openxmlformats-officedocument."
    "spreadsheetml.sheet"
)


def sample_xlsx():
    """Genera un Excel válido con datos simulados."""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ventas"

    sheet.append([
        "ID Venta",
        "Cliente",
        "Total",
    ])

    sheet.append([
        1,
        "Cliente Prueba",
        100.50,
    ])

    output = BytesIO()
    workbook.save(output)
    workbook.close()

    return output.getvalue()


class SalesXlsxApiTests(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app)

    @classmethod
    def tearDownClass(cls):
        cls.client.close()

    def setUp(self):
        app.dependency_overrides.clear()
        app.dependency_overrides[get_db] = lambda: object()

    def tearDown(self):
        app.dependency_overrides.clear()

    def authenticate_as(self, role):
        app.dependency_overrides[get_current_user] = (
            lambda: SimpleNamespace(id=1, role=role)
        )

    def mock_export(self):
        report_patch = patch(
            "app.api.routes.reports.build_sales_report",
            return_value=object(),
        )

        xlsx_patch = patch(
            "app.api.routes.reports.build_sales_xlsx",
            return_value=sample_xlsx(),
        )

        return report_patch, xlsx_patch

    def test_administrator_allowed(self):
        """Administrador puede exportar Excel."""

        self.authenticate_as(ROLE_ADMIN)
        report_patch, xlsx_patch = self.mock_export()

        with report_patch, xlsx_patch:
            response = self.client.get(URL)

        self.assertEqual(response.status_code, 200)

    def test_manager_allowed(self):
        """Gerente puede exportar Excel."""

        self.authenticate_as(ROLE_MANAGER)
        report_patch, xlsx_patch = self.mock_export()

        with report_patch, xlsx_patch:
            response = self.client.get(URL)

        self.assertEqual(response.status_code, 200)

    def test_unauthenticated(self):
        """Sin autenticación devuelve HTTP 401."""

        response = self.client.get(URL)

        self.assertEqual(response.status_code, 401)

    def test_unauthorized_roles(self):
        """Los roles no autorizados reciben HTTP 403."""

        for role in (
            ROLE_ANALYST,
            ROLE_SELLER,
            ROLE_WAREHOUSE,
        ):
            with self.subTest(role=role):
                self.authenticate_as(role)

                response = self.client.get(URL)

                self.assertEqual(response.status_code, 403)

    def test_content_type(self):
        """La respuesta utiliza el MIME de Excel."""

        self.authenticate_as(ROLE_MANAGER)
        report_patch, xlsx_patch = self.mock_export()

        with report_patch, xlsx_patch:
            response = self.client.get(URL)

        self.assertEqual(
            response.headers["content-type"],
            XLSX_MIME,
        )

    def test_download_filename(self):
        """El archivo incluye el período seleccionado."""

        self.authenticate_as(ROLE_MANAGER)
        report_patch, xlsx_patch = self.mock_export()

        with report_patch, xlsx_patch:
            response = self.client.get(URL)

        disposition = response.headers[
            "content-disposition"
        ]

        self.assertIn("attachment", disposition)
        self.assertIn(
            "sales_report_2026-10-01_2026-10-31.xlsx",
            disposition,
        )

    def test_no_store(self):
        """El navegador no almacena el Excel en caché."""

        self.authenticate_as(ROLE_MANAGER)
        report_patch, xlsx_patch = self.mock_export()

        with report_patch, xlsx_patch:
            response = self.client.get(URL)

        self.assertEqual(
            response.headers["cache-control"],
            "no-store",
        )

    def test_invalid_dates(self):
        """Rechaza intervalos de fechas invertidos."""

        self.authenticate_as(ROLE_MANAGER)

        with patch(
            "app.api.routes.reports.build_sales_report"
        ) as service:
            response = self.client.get(
                "/api/reports/sales/export/xlsx"
                "?date_from=2026-10-31"
                "&date_to=2026-10-01"
            )

        self.assertEqual(response.status_code, 422)
        service.assert_not_called()

    def test_invalid_filters(self):
        """Rechaza identificadores inválidos."""

        self.authenticate_as(ROLE_MANAGER)

        for parameter in (
            "branch_id=0",
            "seller_id=-1",
        ):
            with self.subTest(parameter=parameter):
                response = self.client.get(
                    URL + "&" + parameter
                )

                self.assertEqual(response.status_code, 422)

    def test_filters_forwarded(self):
        """Transmite sucursal y vendedor al servicio."""

        self.authenticate_as(ROLE_MANAGER)
        report_patch, xlsx_patch = self.mock_export()

        with report_patch as report_service, xlsx_patch:
            response = self.client.get(
                URL
                + "&branch_id=1"
                + "&seller_id=2"
            )

        self.assertEqual(response.status_code, 200)

        filters = report_service.call_args.kwargs[
            "filters"
        ]

        self.assertEqual(filters.branch_id, 1)
        self.assertEqual(filters.seller_id, 2)

    def test_excel_readable(self):
        """El archivo descargado puede abrirse con openpyxl."""

        self.authenticate_as(ROLE_MANAGER)
        report_patch, xlsx_patch = self.mock_export()

        with report_patch, xlsx_patch:
            response = self.client.get(URL)

        workbook = load_workbook(
            BytesIO(response.content),
            data_only=True,
        )

        try:
            self.assertEqual(
                workbook.sheetnames,
                ["Ventas"],
            )

            sheet = workbook["Ventas"]

            self.assertEqual(
                sheet["A2"].value,
                1,
            )
            self.assertEqual(
                sheet["C2"].value,
                100.50,
            )
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
